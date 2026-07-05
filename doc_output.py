# -*- coding: utf-8 -*-
"""
Saves generated content into Word documents instead of dumping markdown to
the terminal, since John reviews and copies posts from a docx, not a console.

If the target doc is open in Word (Windows locks it for writing), the save
falls back to a sibling "(unsaved - original was open)" file so a draft that
already cost Gemini credits is never lost -- important for the unattended
weekly scheduler, where John may have the review doc open when it fires.
"""

import os
from datetime import date
from docx import Document


def _fallback_path(doc_path: str) -> str:
    """Sibling filename used when the real target is locked open in Word."""
    root, ext = os.path.splitext(doc_path)
    return root + " (unsaved - original was open)" + ext


def _open_or_create(doc_path: str):
    """Open an existing doc, or create a new one titled after its filename."""
    if os.path.exists(doc_path):
        return Document(doc_path)
    doc = Document()
    title = os.path.splitext(os.path.basename(doc_path))[0]
    doc.add_heading(title, level=1)
    return doc


def _add_entry(doc, heading: str, body: str) -> None:
    """Append one dated H2 section plus the body paragraphs."""
    doc.add_heading(f"{heading} -- {date.today().isoformat()}", level=2)
    for line in body.split("\n"):
        doc.add_paragraph(line)


def append_to_doc(doc_path: str, heading: str, body: str) -> str:
    """Append a dated heading and body to doc_path, creating the file (with a
    title matching its filename) if it does not exist yet. Returns the path
    actually written. If doc_path is locked (open in Word), the draft is saved
    to a fallback sibling file instead of being lost; only if that also fails
    does it stop with a plain-English error."""
    doc = _open_or_create(doc_path)
    _add_entry(doc, heading, body)
    try:
        doc.save(doc_path)
        return doc_path
    except PermissionError:
        # Original is open in Word. Do not throw the (already-paid-for) draft
        # away - accumulate it in a fallback doc and keep the run going.
        fallback = _fallback_path(doc_path)
        fdoc = _open_or_create(fallback)
        _add_entry(fdoc, heading, body)
        try:
            fdoc.save(fallback)
        except PermissionError:
            raise SystemExit(
                f"\n[DOC] Could not save '{doc_path}' or its fallback "
                f"'{fallback}' - both look locked. Close them in Word, then "
                "run again.\n"
            )
        print(
            f"[DOC] '{doc_path}' was open in Word, so this draft was saved to "
            f"'{fallback}' instead (nothing lost). Close the original and "
            "future drafts will save to it again."
        )
        return fallback
